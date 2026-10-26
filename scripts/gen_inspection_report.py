#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""卫星应急车检测报告参数化生成脚本。

用法：
    python3 gen_inspection_report.py <input.json> <output.docx>

input.json 结构见下方 SAMPLE 或 README 注释；由 Agent 从用户提供的检测清单（Excel/表格/图片/文字）
提取数据后构造。报告骨架、判定分类统计、附图占位符自动生成，严格遵循 references/standards.md
的结构与占位规范。
"""
import json
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SAMPLE = {
    "vehicle": {
        "report_no": "____________", "project": "云南移动卫星应急车损坏评估检测",
        "plate": "云A6J6F2", "vehicle_type": "动中通卫星应急通信车",
        "nature": "受冲击损坏初步评估检测", "location": "云南省____________",
        "date": "2026年10月21日", "unit": "____________",
        "env_temp": "____ ℃", "env_humidity": "____ %RH",
        "env_weather": "____", "env_site": "____（空旷无遮挡场地，天线架设区域平整）",
        "overview": "采用Ku频段车载动中通VSAT卫星通信系统，可实现行车及驻车状态下的卫星链路建立与语音、视频、数据传输，用于突发事件的现场应急通信保障。"
    },
    "items": [
        {"no": "1", "category": "天线系统结构检测", "name": "天线俯仰钢丝", "qty": "12根",
         "standard": "结构件完好、无断裂", "actual": "全部断裂", "verdict": "严重损伤，需更换",
         "method": "外观结构检查", "detail": "共12根，实测全部断裂，天线俯仰角度支撑及回中功能丧失。"},
        {"no": "2", "category": "射频链路部件检测", "name": "KU-BUC（功放）", "qty": "1台",
         "standard": "输出频率14.0～14.5GHz；频率稳定度优于±10kHz；功率：饱和功率/1dB压缩点",
         "actual": "待检测（受严重冲击）", "verdict": "待进一步检测",
         "method": "外观检查；受严重冲击暂不满足通电条件，指标检测（JT/T 1007.3—2024表6）待返厂实施",
         "detail": "受严重冲击，是否损伤待检测。"}
    ]
}


def set_font(style, cn, en='Times New Roman', size=None, bold=None, color=RGBColor(0, 0, 0)):
    style.font.name = en
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn('w:ascii'), en)
    rfonts.set(qn('w:hAnsi'), en)
    rfonts.set(qn('w:eastAsia'), cn)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def build(data, out_path):
    v = data['vehicle']
    items = data['items']
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

    normal = doc.styles['Normal']
    set_font(normal, '宋体', size=12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)

    tstyle = doc.styles['Title']
    set_font(tstyle, '黑体', size=18, bold=True)
    tstyle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tstyle.paragraph_format.space_after = Pt(10)
    tstyle.paragraph_format.first_line_indent = Pt(0)

    for lv, sz, before, after in ((1, 16, 14, 6), (2, 14, 11, 5), (3, 12, 9, 4)):
        h = doc.styles['Heading %d' % lv]
        set_font(h, '黑体', size=sz, bold=True)
        h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.first_line_indent = Pt(0)

    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    run.font.size = Pt(10.5)

    def body(text, indent=True, align=None, bold=False):
        p = doc.add_paragraph()
        r = p.add_run(text); r.font.bold = bold
        if not indent:
            p.paragraph_format.first_line_indent = Pt(0)
        if align == 'center':
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'left':
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    def heading(text, level):
        p = doc.add_heading('', level=level)
        p.add_run(text).font.bold = True
        return p

    def photo_placeholder():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.add_run('【附件：现场检测照片____（此处自行插入图片）】').font.bold = True

    def set_cell(cell, text, bold=False, align='left', size=10.5):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold
        p.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER,
                       'right': WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.LEFT)

    def style_table(tbl):
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '000000')
            borders.append(el)
        tbl._tbl.tblPr.append(borders)

    def set_row_cant_split(row, header=False):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if header:
            trPr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(OxmlElement('w:cantSplit'))

    def shade_cell(cell, fill='D9D9D9'):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

    def kv_table(rows, w1=4.2, w2=11.8):
        t = doc.add_table(rows=len(rows), cols=2)
        style_table(t)
        for i, (k, vv) in enumerate(rows):
            set_cell(t.rows[i].cells[0], k, bold=True, align='center')
            set_cell(t.rows[i].cells[1], vv, align='left')
            shade_cell(t.rows[i].cells[0], 'F2F2F2')
            set_row_cant_split(t.rows[i])
        for row in t.rows:
            row.cells[0].width = Cm(w1); row.cells[1].width = Cm(w2)
        return t

    def data_table(headers, rows, widths, cell_aligns=None):
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        style_table(t)
        for j, hh in enumerate(headers):
            set_cell(t.rows[0].cells[j], hh, bold=True, align='center')
            shade_cell(t.rows[0].cells[j], 'D9D9D9')
        set_row_cant_split(t.rows[0], header=True)
        for i, row in enumerate(rows, start=1):
            for j, vv in enumerate(row):
                a = (cell_aligns[j] if cell_aligns else 'center')
                set_cell(t.rows[i].cells[j], vv, align=a)
            set_row_cant_split(t.rows[i])
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
        return t

    def caption(text, above=False):
        cap = doc.add_paragraph()
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Pt(0)
        if above:
            cap.paragraph_format.space_after = Pt(4)
        else:
            cap.paragraph_format.space_before = Pt(4)
        cap.add_run(text).font.size = Pt(10.5)

    # ---------- 封面 ----------
    tp = doc.add_paragraph('', style='Title')
    tp.add_run('卫星通信应急车损坏评估检测报告').font.bold = True
    body('（检测对象：%s %s）' % (v['plate'], v['vehicle_type']), indent=False, align='center', bold=True)
    body('文件编号：%s    版本号：V____    密级：□公开 □内部 □秘密' % v['report_no'], indent=False, align='center')
    doc.add_paragraph()

    # ---------- 1 报告基本信息 ----------
    heading('1 报告基本信息', 1)
    kv_table([
        ('报告编号', v['report_no']), ('项目名称', v['project']),
        ('检测对象', '%s卫星应急车（车载动中通VSAT卫星应急通信车）' % v['plate']),
        ('检测性质', v['nature']), ('检测地点', v['location']),
        ('检测时间', v['date']), ('检测单位', v['unit']),
        ('编制人/检测人', '____________'), ('审核人', '____________'),
    ])

    # ---------- 2 检测概述 ----------
    heading('2 检测概述', 1)
    heading('2.1 检测目的', 2)
    body('对%s卫星应急车因外力冲击造成的天馈及射频部件损伤情况进行初步评估检测，查明各部件损伤状态，判定是否具备继续使用条件，为后续更换、修复、返厂检测及整机恢复提供依据。' % v['plate'])
    heading('2.2 检测范围', 2)
    body('本次检测范围覆盖天线系统、射频链路、机械驱动系统及辅助结构四类部件，共%d个检测项。车辆底盘、卫星Modem、供电系统及整车电气部分未列入本次损坏评估范围。' % len(items))
    heading('2.3 检测依据', 2)
    body('本次检测及评估判定依据下列现行有效标准及规范执行：')
    data_table(
        ['序号', '标准编号', '标准名称', '适用说明'],
        [('1', 'YD/T 2675-2015', '《应急通信车载卫星地球站通用技术要求》', '车载卫星地球站总体技术要求'),
         ('2', 'JT/T 1007.3—2024', '《交通移动应急通信指挥系统 第3部分：设备检测规范》', '检测项目、量化指标及检测方法（表4/6/7/8）'),
         ('3', 'GB/T 51266-2018', '《应急通信系统工程技术标准》', '应急通信系统工程建设与验收通用要求')],
        [1.4, 3.6, 7.6, 3.4], cell_aligns=['center', 'center', 'left', 'left'])

    # ---------- 3 车辆及设备概况 ----------
    heading('3 车辆及设备概况', 1)
    heading('3.1 车辆总体情况', 2)
    body('%s卫星应急车为移动卫星通信应急保障车辆，%s本车近期遭受外力冲击，天线及射频部件出现多处损伤，本次检测针对上述损伤开展初步评估。' % (v['plate'], v['overview']))
    heading('3.2 设备组成', 2)
    cats = []
    for it in items:
        if it['category'] not in cats:
            cats.append(it['category'])
    rows = []
    for c in cats:
        names = [it['name'] for it in items if it['category'] == c]
        rows.append((c, '；'.join(names), str(len([it for it in items if it['category'] == c])), '见分项检测'))
    body('本车卫星通信系统主要设备组成如下：')
    data_table(['系统', '主要设备', '数量', '备注'], rows, [3.0, 8.4, 1.6, 3.0],
               cell_aligns=['center', 'left', 'center', 'left'])
    body('注：车辆底盘、供电系统及卫星Modem技术资料待补充，本报告未对其状态作出评定。')

    # ---------- 4 检测条件与检测设备 ----------
    heading('4 检测条件与检测设备', 1)
    heading('4.1 检测环境条件', 2)
    body('本次外观及结构检测在露天场地进行，环境条件记录如下（现场实测后填写）：')
    kv_table([('环境温度', v['env_temp']), ('相对湿度', v['env_humidity']),
              ('天气状况', v['env_weather']), ('检测场地', v['env_site'])])
    heading('4.2 检测仪器与器具', 2)
    body('本报告涉及的检测仪器设备应符合JT/T 1007.3—2024第5.2.2条对检测设备的基本技术要求。拟用及已用检测器具如下：')
    data_table(
        ['序号', '仪器/器具名称', '主要技术指标要求', '用途'],
        [('1', '频谱仪', '频率范围100kHz～3GHz；分辨率带宽100kHz～1MHz；视频带宽10Hz～1MHz；显示电平误差<1.5dB', '射频链路指标专项检测（后续实施）'),
         ('2', '通用外观检查与测量器具', '游标卡尺、卷尺、角度测量工具、扭矩工具等', '结构形变、间隙及尺寸检查'),
         ('3', '倾角仪/水平仪', '分辨率不低于0.1°', '天线姿态及限位位置检查'),
         ('4', '万用表', '常规电气参数测量', '供电及控制回路状态核查')],
        [1.4, 4.6, 6.4, 3.6], cell_aligns=['center', 'left', 'left', 'left'])
    body('注：上述仪器仪表应处于检定/校准有效期内，检测前进行功能确认。')
    heading('4.3 检测安全注意事项', 2)
    body('（1）KU-BUC、KU-LNB等有源器件受严重冲击，在未排除内部损伤前严禁盲目加电测试，防止二次损坏及人身伤害；')
    body('（2）天线反射面板多处破损，检查及搬运作业时应佩戴防护手套，防止金属毛刺割伤；')
    body('（3）拆解机械驱动部件前应切断天线控制系统供电，并采取防坠落、防倾覆措施；')
    body('（4）吊装及高处作业应符合相关安全操作规程。')

    # ---------- 5 检测方法与判定准则 ----------
    heading('5 检测方法与判定准则', 1)
    heading('5.1 检测方法', 2)
    body('本次初步评估检测采用外观结构检查法为主、出厂检测资料查验为辅的检测方式，各检测项的具体方法说明如下：')
    body('（1）外观结构检查法：对部件外观状态、形变、断裂、间隙进行目视及量具检查，如实记录损伤状况；')
    body('（2）出厂检测报告查验法：对有出厂量化指标的部件（天线、功放、LNB等），查验出厂检测报告并作为标准要求值的核对依据（依据JT/T 1007.3—2024表4、表6、表7）；')
    body('（3）专项测试法：对具备通电及加电条件的部件，依据JT/T 1007.3—2024第5.2.3～5.2.6条规定的检测方法实施射频指标专项检测；本次受部件损伤状态限制，有源器件及转动部件的专项测试待拆解、返厂或修复后实施。')
    heading('5.2 判定准则', 2)
    body('按部件损伤状态，检测判定结论分为以下三类：')
    body('（1）严重损伤：部件已丧失功能或无法修复，判定为需更换；')
    body('（2）受损可修复：部件存在损伤但具备修复条件，判定为需修复处理；')
    body('（3）待进一步检测：部件处于受冲击部位，或需拆解、通电后方可确定损伤程度，判定为待检。')
    body('有源器件（KU-BUC、KU-LNB）受严重冲击，依据YD/T 2675-2015有关要求及JT/T 1007.3—2024第5.2.4、5.2.5条检测条件，在未排除内部损伤前不具备通电测试条件，按待进一步检测处理。')

    # ---------- 6 分项检测结果 ----------
    heading('6 分项检测结果', 1)
    sec_no = 1
    for c in cats:
        heading('6.%d %s' % (sec_no, c), 2)
        sec_no += 1
        for it in [x for x in items if x['category'] == c]:
            heading('6.%d.%s %s' % (sec_no - 1, it['no'], it['name']), 3)
            body('检测方法：%s。' % it['method'], indent=False)
            body('检测结果：%s' % it['detail'])
            body('判定：%s。' % it['verdict'], indent=False, bold=True)
            photo_placeholder()

    # ---------- 7 检测结果汇总表 ----------
    heading('7 检测结果汇总表', 1)
    caption('表7-1 %s卫星应急车部件损坏检测结果汇总表' % v['plate'], above=True)
    sum_rows = [(it['no'], it['name'], it['qty'], it['standard'], it['actual'], it['verdict']) for it in items]
    data_table(['序号', '部件名称', '数量', '标准要求值', '实测状况', '判定结论'], sum_rows,
               [1.2, 3.0, 1.4, 4.8, 3.0, 2.6],
               cell_aligns=['center', 'left', 'center', 'left', 'left', 'left'])
    body('注：标准要求值引自JT/T 1007.3—2024《交通移动应急通信指挥系统 第3部分：设备检测规范》表4、表6、表7及YD/T 2675-2015相关要求；f为工作频率（GHz）。', indent=False)

    # ---------- 8 综合评定（自动统计） ----------
    heading('8 综合评定', 1)
    body('综合本次检测，%s卫星应急车因外力冲击，天线系统、射频链路、机械驱动系统及辅助结构均存在不同程度损伤，评定如下：' % v['plate'])
    heavy = [it['name'] for it in items if it['verdict'].startswith('严重损伤')]
    repair = [it['name'] for it in items if it['verdict'].startswith('受损可修复')]
    pending = [it['name'] for it in items if it['verdict'].startswith('待进一步')]
    if heavy:
        body('（1）%s共%d类部件判定为严重损伤，其功能及结构完整性已不能满足使用要求，须更换处理；' % ('、'.join(heavy), len(heavy)))
    if pending:
        body('（2）%s等部件处于受冲击部位或需专项检测，须拆解/返厂后方可最终判定；' % '、'.join(pending))
    if repair:
        body('（3）%s存在损伤，具备修复条件，按原工艺修复并补强后可继续使用。' % '、'.join(repair))
    body('风险等级评定：本车天馈及射频系统损伤程度高，当前不具备入网通信及应急保障能力，评定为高风险设备状态，必须完成更换、修复并经整机复测合格后方可恢复执行应急通信保障任务。')

    # ---------- 9 现存问题及整改建议 ----------
    heading('9 现存问题说明及运维整改建议', 1)
    heading('9.1 现存问题', 2)
    body('（1）无源天馈部件多处达到报废标准：%s；' % ('、'.join(heavy) if heavy else '待确认'))
    body('（2）有源器件（KU-BUC、KU-LNB）损伤状态未知，存在盲目加电导致二次损坏的风险；')
    body('（3）关键传动与转动部件损伤状态未定，须拆解检测；')
    body('（4）维修期间该车无法承担应急通信保障任务，存在应急保障能力空缺。')
    heading('9.2 运维整改建议', 2)
    if heavy:
        body('（1）对判定为严重损伤的部件实施更换：%s；' % '、'.join(heavy))
    body('（2）对判定为待进一步检测的部件组织拆解/返厂检测，确认损伤程度后再行处置；')
    body('（3）KU-BUC、KU-LNB在未确认内部损伤状态前禁止盲目加电，送返厂检测或直接更换；')
    body('（4）对受损可修复部件按原工艺修复并补强；')
    body('（5）全部更换、修复完成后，依据YD/T 2675-2015《应急通信车载卫星地球站通用技术要求》及JT/T 1007.3—2024第5.2.3～5.2.6条，开展整机性能复测，包括：天线对星精度、发射/接收增益、G/T值、驻波比、收发端口隔离度、上下行链路指标、动中通动态跟踪能力（车速80km/h工况下指向精度）、语音/视频/数据业务联通性及电源稳定性等，复测合格后方可恢复入网；')
    body('（6）修复期间该车不得执行应急通信保障任务，建议同步推进事故责任认定与保险理赔评估，并统筹调度其他应急资源填补保障空缺。')

    # ---------- 10 签署栏 ----------
    heading('10 签署栏', 1)
    body('以下签署经确认后生效：', indent=False)
    kv_table([('检测人（签字）', '____________'), ('审核人（签字）', '____________'),
              ('批准人（签字）', '____________'), ('检测单位（盖章）', '____________')])

    # ---------- 附录A ----------
    doc.add_page_break()
    heading('附录A 缩略语', 1)
    data_table(['缩略语', '英文全称', '中文含义'],
               [('BUC', 'Block Up-Converter', '上变频功率放大器'),
                ('LNB', 'Low Noise Block', '低噪声下变频器'),
                ('VSAT', 'Very Small Aperture Terminal', '甚小口径终端'),
                ('G/T', 'Gain-to-Noise Temperature Ratio', '接收系统品质因数'),
                ('EIRP', 'Equivalent Isotropically Radiated Power', '等效全向辐射功率'),
                ('VSWR', 'Voltage Standing Wave Ratio', '电压驻波比'),
                ('GPS', 'Global Positioning System', '全球定位系统'),
                ('Ku频段', 'Ku-band', '12～18GHz频段（上行约14GHz、下行约12GHz）'),
                ('Modem', 'Modulator and Demodulator', '调制解调器')],
               [2.6, 8.2, 5.2], cell_aligns=['center', 'left', 'left'])
    body('')
    body('（本报告共____页，正文及附表未尽事宜按检测依据标准执行。）', indent=False)

    doc.save(out_path)
    print('saved:', out_path)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1], 'r', encoding='utf-8') as f:
        data = json.load(f)
    build(data, sys.argv[2])
